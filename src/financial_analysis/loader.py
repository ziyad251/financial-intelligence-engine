"""Load monthly financial-statement worksheets into a normalised row/group tree.

The loader makes no assumptions about account names. It finds the header row by
looking for month columns, then walks the rows top to bottom:

* a labelled row with no values opens a group (a section header such as "Income");
* a row labelled "Total <group>" closes the innermost matching open group and is
  recorded as that group's reported subtotal — formula or hard-coded alike;
* rows named like a derived line (Gross Profit, Net Income, ...) and any other row
  whose cells are formulas referencing other cells are recorded as reported,
  calculated rows;
* everything else is a leaf account — the only rows the pipeline ever sums.

Keeping subtotals and calculated rows out of the leaf set is what prevents
double-counting: every figure the pipeline calculates is a sum of leaves, and the
workbook's own subtotals are used only to reconcile against.
"""

from __future__ import annotations

import hashlib
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .config import Config, label_in, normalize_label
from .formula import FormulaError, FormulaEvaluator, references_cells
from .issues import ERROR, INFO, WARNING, AnalysisError, Issue
from .periods import Period, consecutive_gaps, parse_period

HEADER_SCAN_ROWS = 30

LEAF = "leaf"
SUBTOTAL = "subtotal"
DERIVED = "derived"

Series = dict[Period, Decimal | None]


@dataclass
class Row:
    sheet_row: int
    label: str
    kind: str
    path: tuple[str, ...] = ()  # labels of the enclosing groups, outermost first
    values: Series = field(default_factory=dict)
    cells: dict[Period, str] = field(default_factory=dict)  # period -> cell coordinate
    has_formula: bool = False
    # Periods whose value Excel had not cached, so the pipeline recalculated it from the cell's formula.
    recalculated: list[Period] = field(default_factory=list)

    @property
    def path_label(self) -> str:
        return " > ".join((*self.path, self.label))


@dataclass
class Group:
    label: str
    indent: int
    path: tuple[str, ...]  # ancestors, outermost first
    header_row: int
    rows: list[Row] = field(default_factory=list)  # direct leaf rows
    children: list["Group"] = field(default_factory=list)
    total: Row | None = None
    closed_explicitly: bool = False

    def leaves(self) -> list[Row]:
        """All leaf rows under this group, in sheet order, each exactly once."""
        out = list(self.rows)
        for child in self.children:
            out.extend(child.leaves())
        return sorted(out, key=lambda r: r.sheet_row)

    def walk(self) -> Iterator["Group"]:
        yield self
        for child in self.children:
            yield from child.walk()


@dataclass
class Statement:
    sheet: str
    title_lines: list[str]
    periods: list[Period]  # ascending
    rows: list[Row]
    groups: list[Group]  # top level
    orphan_leaves: list[Row]

    def all_groups(self) -> list[Group]:
        return [g for top in self.groups for g in top.walk()]

    def leaves(self) -> list[Row]:
        return [r for r in self.rows if r.kind == LEAF]

    def find_groups(self, aliases: list[str]) -> list[Group]:
        """Groups whose label matches one of the aliases, outermost first."""
        found = [g for g in self.all_groups() if label_in(g.label, aliases)]
        return sorted(found, key=lambda g: (len(g.path), g.header_row))

    def find_rows(self, aliases: list[str], kinds: tuple[str, ...]) -> list[Row]:
        return [r for r in self.rows if r.kind in kinds and label_in(r.label, aliases)]


@dataclass
class LoadedWorkbook:
    path: Path
    sha256: str
    sheet_names: list[str]
    income_statement: Statement
    balance_sheet: Statement | None
    issues: list[Issue]


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _find_sheet(names: list[str], aliases: list[str]) -> str | None:
    for alias in aliases:
        for name in names:
            if normalize_label(name) == normalize_label(alias):
                return name
    return None


def _find_header(ws) -> tuple[int, dict[Period, int], list[tuple[int, Period]]]:
    """Return (header row, {period: column}, duplicate (column, period) pairs)."""
    for r in range(1, min(HEADER_SCAN_ROWS, ws.max_row) + 1):
        parsed = []
        for c in range(1, ws.max_column + 1):
            period = parse_period(ws.cell(r, c).value)
            if period is not None:
                parsed.append((c, period))
        if len(parsed) >= 2:
            columns: dict[Period, int] = {}
            duplicates = []
            for c, period in parsed:
                if period in columns:
                    duplicates.append((c, period))
                else:
                    columns[period] = c
            return r, columns, duplicates
    raise AnalysisError(f"Sheet '{ws.title}': no header row with at least two month columns in the first {HEADER_SCAN_ROWS} rows")


class _SheetParser:
    def __init__(self, ws_formulas, ws_values, config: Config, statement_key: str, issues: list[Issue]):
        self.wf, self.wv = ws_formulas, ws_values
        self.config = config
        self.statement_key = statement_key
        self.issues = issues
        self.sheet = ws_formulas.title
        self.evaluator = FormulaEvaluator(lambda col, row: self.wf.cell(row, col).value)
        derived = config.income_statement["derived_rows"] if statement_key == "income_statement" else {}
        self.derived_aliases = [alias for names in derived.values() for alias in names]
        self.evaluated: dict[int, list[tuple[Period, str]]] = {}  # sheet row -> [(period, cell)] recalculated

    def issue(self, code: str, severity: str, message: str, **kw) -> None:
        self.issues.append(Issue(code, severity, message, sheet=self.sheet, **kw))

    def cell_value(self, row: int, col: int, period: Period, label: str) -> tuple[Decimal | None, bool]:
        """Return (numeric value or None, is_formula). Problems are recorded as issues."""
        raw = self.wf.cell(row, col).value
        coord = f"{get_column_letter(col)}{row}"
        if _blank(raw):
            return None, False
        if isinstance(raw, bool) or isinstance(raw, (datetime, date)):
            self.issue("non_numeric_value", WARNING, f"{coord} holds {raw!r}, not a number; treated as missing",
                       account=label, period=period.key)
            return None, False
        if isinstance(raw, (int, float)):
            return Decimal(str(raw)), False
        text = str(raw).strip()
        if not text.startswith("="):
            self.issue("non_numeric_value", WARNING, f"{coord} holds text {text!r}, not a number; treated as missing",
                       account=label, period=period.key)
            return None, False
        cached = self.wv.cell(row, col).value
        if isinstance(cached, (int, float)) and not isinstance(cached, bool):
            return Decimal(str(cached)), True
        if cached is not None:  # e.g. '#REF!', '#DIV/0!'
            self.issue("formula_error", WARNING, f"{coord} formula {text} shows {cached!r}; treated as missing",
                       account=label, period=period.key)
            return None, True
        try:
            value = self.evaluator.evaluate(text)
        except FormulaError as exc:
            self.issue("formula_unevaluated", WARNING,
                       f"{coord} formula {text} has no cached value and could not be evaluated ({exc}); treated as missing",
                       account=label, period=period.key)
            return None, True
        self.evaluated.setdefault(row, []).append((period, coord))  # reported once per row, see _report_recalculated
        return value, True

    def parse(self) -> Statement:
        ws = self.wf
        header_row, columns, duplicates = _find_header(ws)
        for c, period in duplicates:
            self.issue("duplicate_period", ERROR,
                       f"Column {get_column_letter(c)} repeats period {period.label}; only column "
                       f"{get_column_letter(columns[period])} is used", period=period.key)
        first_col = min(columns.values())
        label_cols = range(1, first_col)
        periods = sorted(columns)

        used_cols = set(columns.values())
        for c in range(first_col, ws.max_column + 1):
            if c in used_cols:
                continue
            numeric = [r for r in range(header_row + 1, ws.max_row + 1) if not _blank(ws.cell(r, c).value)]
            if numeric:
                head = ws.cell(header_row, c).value
                self.issue("ignored_column", WARNING,
                           f"Column {get_column_letter(c)} (header {head!r}) has data but is not a single month; ignored",
                           details={"column": get_column_letter(c), "header": head, "non_blank_cells": len(numeric)})

        title_lines = []
        for r in range(1, header_row):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.strip():
                    title_lines.append(v.strip())

        rows: list[Row] = []
        top: list[Group] = []
        orphans: list[Row] = []
        stack: list[Group] = []

        def close_implicitly(group: Group, reason: str) -> None:
            self.issue("group_without_total", WARNING,
                       f"Group '{group.label}' (row {group.header_row}) has no 'Total {group.label}' row ({reason}); "
                       "its accounts are still used but it cannot be reconciled", account=group.label)

        for r in range(header_row + 1, ws.max_row + 1):
            label, indent = None, 0
            for c in label_cols:
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.strip():
                    label, indent = v.strip(), c - label_cols.start
                    break
            raw_present = any(not _blank(ws.cell(r, c).value) for c in columns.values())
            if label is None:
                if raw_present:
                    self.issue("unlabelled_row", WARNING, f"Row {r} has values but no account label; ignored")
                continue

            if not raw_present:
                group = Group(label, indent, tuple(g.label for g in stack), r)
                (stack[-1].children if stack else top).append(group)
                stack.append(group)
                continue

            row = Row(r, label, LEAF, tuple(g.label for g in stack))
            for period, c in columns.items():
                row.values[period], is_formula = self.cell_value(r, c, period, label)
                row.cells[period] = f"{get_column_letter(c)}{r}"
                raw = ws.cell(r, c).value
                if is_formula and isinstance(raw, str) and references_cells(raw):
                    row.has_formula = True
            rows.append(row)

            norm = normalize_label(label)
            if norm.startswith("total "):
                target = norm[len("total "):]
                idx = next((i for i in range(len(stack) - 1, -1, -1) if normalize_label(stack[i].label) == target), None)
                row.kind = SUBTOTAL
                if idx is None:
                    self.issue("unmatched_total", WARNING,
                               f"Row {r} '{label}' does not close any open group; excluded from account sums",
                               account=label)
                    continue
                for inner in stack[idx + 1:]:
                    close_implicitly(inner, f"closed by '{label}'")
                group = stack[idx]
                group.total, group.closed_explicitly = row, True
                row.path = group.path
                del stack[idx:]
                continue

            if label_in(label, self.derived_aliases):
                row.kind = DERIVED
                continue
            if row.has_formula:
                row.kind = DERIVED
                self.issue("formula_row_excluded", INFO,
                           f"Row {r} '{label}' is calculated from other cells, so it is treated as a reported "
                           "calculated line and not summed as an account", account=label)
                continue

            if stack:
                stack[-1].rows.append(row)
            else:
                orphans.append(row)
                self.issue("orphan_row", WARNING,
                           f"Row {r} '{label}' is not inside any section; excluded from section totals", account=label)

        for group in stack:
            close_implicitly(group, "sheet ended")

        self._apply_blank_policy(rows)
        self._report_recalculated(rows)
        for gap in consecutive_gaps(periods):
            self.issue("missing_period", WARNING,
                       f"{gap.label} is missing between the first and last months; comparisons that need it are not calculated",
                       period=gap.key)
        return Statement(self.sheet, title_lines, periods, rows, top, orphans)

    def _report_recalculated(self, rows: list[Row]) -> None:
        """One info issue per row whose formula cells had no value cached by Excel (the file was written by a
        tool that does not calculate, or saved without recalculating). The values are the pipeline's own
        evaluation of the formulas, not figures Excel reported, and the issue says what that affects."""
        for row in rows:
            cells = self.evaluated.get(row.sheet_row)
            if not cells:
                continue
            row.recalculated = sorted(p for p, _ in cells)
            coords = [c for _, c in sorted(cells)]
            if row.kind == LEAF:
                use = "This account's recalculated values are used in the calculated figures."
            else:
                use = ("This is a workbook total, used only to reconcile against, so its check compares the "
                       "pipeline's sums with the workbook's formulas, not with Excel's saved results.")
            span = coords[0] if len(coords) == 1 else f"{coords[0]}..{coords[-1]}"
            self.issue("formula_evaluated", INFO,
                       f"Row {row.sheet_row} '{row.label}': {len(coords)} formula cell(s) ({span}) had no value cached "
                       f"by Excel and were recalculated by the pipeline. {use}",
                       account=row.label, details={"row": row.sheet_row, "row_kind": row.kind, "cells": coords,
                                                   "periods": [p.key for p in row.recalculated]})

    def _apply_blank_policy(self, rows: list[Row]) -> None:
        blanks_as_zero = self.config.blank_cells_as_zero
        filled = 0
        for row in rows:
            if row.kind != LEAF:
                continue
            missing = [p for p, v in row.values.items() if v is None and _blank(self.wf[row.cells[p]].value)]
            if not missing:
                continue
            if blanks_as_zero:
                for p in missing:
                    row.values[p] = Decimal(0)
                filled += len(missing)
            else:
                self.issue("missing_value", WARNING,
                           f"'{row.path_label}' is blank for {', '.join(p.label for p in sorted(missing))}; "
                           "totals that include it are not calculated for those months",
                           account=row.label, details={"periods": [p.key for p in sorted(missing)]})
        if filled:
            self.issue("blank_as_zero", INFO,
                       f"{filled} blank account cells treated as zero (configured: blank_cells_as_zero)")


def load_workbook(path: str | Path, config: Config) -> LoadedWorkbook:
    path = Path(path)
    if not path.is_file():
        raise AnalysisError(f"Workbook not found: {path}")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    try:
        wb_f = openpyxl.load_workbook(path, data_only=False)
        wb_v = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # corrupt, password-protected, .xls/.csv renamed, truncated download, ...
        raise AnalysisError(f"Could not open workbook {path.name}: it is not a readable .xlsx/.xlsm file "
                            f"({type(exc).__name__}: {exc}). Re-export it from Excel or the accounting system.") from exc
    names = wb_f.sheetnames
    issues: list[Issue] = []

    is_name = _find_sheet(names, config.income_statement["sheet_names"])
    if is_name is None:
        raise AnalysisError(f"No income statement sheet found. Sheets: {names}. "
                            "Add the sheet name to income_statement.sheet_names in the company config.")
    income = _SheetParser(wb_f[is_name], wb_v[is_name], config, "income_statement", issues).parse()

    bs_name = _find_sheet(names, config.balance_sheet["sheet_names"])
    balance = None
    if bs_name is None:
        issues.append(Issue("balance_sheet_missing", INFO,
                            "No balance sheet sheet found; balance-sheet checks and metrics are skipped"))
    else:
        balance = _SheetParser(wb_f[bs_name], wb_v[bs_name], config, "balance_sheet", issues).parse()
        if set(balance.periods) != set(income.periods):
            only_is = sorted(set(income.periods) - set(balance.periods))
            only_bs = sorted(set(balance.periods) - set(income.periods))
            issues.append(Issue("period_coverage_mismatch", INFO,
                                "Income statement and balance sheet cover different months",
                                details={"income_statement_only": [p.key for p in only_is],
                                         "balance_sheet_only": [p.key for p in only_bs]}))

    for name in names:
        if name not in (is_name, bs_name):
            issues.append(Issue("sheet_ignored", INFO, f"Sheet '{name}' is not an income statement or balance sheet; ignored",
                                sheet=name))
    return LoadedWorkbook(path, digest, names, income, balance, issues)
