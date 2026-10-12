"""Synthetic workbooks for the golden cases, written with openpyxl (no Excel needed).

Every workbook copies the layout the loader reads in the supplied export:
    row 1 statement title, row 2 company, row 3 period line   (rows 2 = company is how the
    row 5 header: "Account" in column C, one month per column from D   company check reads it)
    row 7 onward: labels in columns A-C (indent 0-2), values from column D

Rows are (indent, label, values):
    values None            -> a group header (a labelled row with no values)
    values list            -> one entry per month: a number, None (blank cell) or text
    values "=SUM({c}8:{c}9)" -> a formula; {c} becomes each month's column letter

openpyxl stores formulas without the values Excel would cache, so the loader recalculates
them (reported as `formula_evaluated`); case 08 checks that this is reported, not hidden.
All companies and figures are fictional. Row numbers in the comments are the sheet rows the
expected results refer to.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

COMPANY = "Golden Test Co"
LABEL_COLS = 3  # labels in A..C, so months start in column D
FIRST_ROW = 7

JAN_MAR_2026 = ["Jan 2026", "Feb 2026", "Mar 2026"]


def write_workbook(path: Path, months: list[str], rows: list[tuple], company: str = COMPANY) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    ws["A1"], ws["A2"], ws["A3"] = "Profit and Loss", company, f"For the period ended {months[-1]}"
    ws.cell(5, LABEL_COLS, "Account")
    for i, month in enumerate(months):
        ws.cell(5, LABEL_COLS + 1 + i, month)
    for r, (indent, label, values) in enumerate(rows, start=FIRST_ROW):
        ws.cell(r, 1 + indent, label)
        if values is None:
            continue
        for i in range(len(months)):
            col = LABEL_COLS + 1 + i
            if isinstance(values, str):
                ws.cell(r, col, values.format(c=get_column_letter(col)))
            elif values[i] is not None:
                ws.cell(r, col, values[i])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def normal_month(path: Path) -> Path:
    """Case 01: three clean months with derived rows (Gross Profit, Operating Income, Net Income)."""
    return write_workbook(path, JAN_MAR_2026, [
        (0, "Income", None),                                        # 7
        (1, "Product Sales", [10000, 11000, 12000]),                # 8
        (1, "Service Revenue", [2000, 2500, 3000]),                 # 9
        (0, "Total Income", "=SUM({c}8:{c}9)"),                     # 10
        (0, "Cost of Goods Sold", None),                            # 11
        (1, "Materials", [4000, 4400, 4800]),                       # 12
        (1, "Freight", [500, 550, 600]),                            # 13
        (0, "Total Cost of Goods Sold", "=SUM({c}12:{c}13)"),       # 14
        (1, "Gross Profit", "={c}10-{c}14"),                        # 15
        (0, "Operating Expenses", None),                            # 16
        (1, "Salaries", [3000, 3000, 3000]),                        # 17
        (1, "Rent", [1000, 1000, 1000]),                            # 18
        (1, "Advertising", [400, 600, 500]),                        # 19
        (1, "Software", [200, 200, 250]),                           # 20
        (1, "Interest Expense", [100, 100, 100]),                   # 21
        (0, "Total Operating Expenses", "=SUM({c}17:{c}21)"),       # 22
        (1, "Operating Income", "={c}15-{c}22"),                    # 23
        (1, "Net Income", "={c}23"),                                # 24
    ])


MOM_TTM_MONTHS = ["Jan 2025", "Feb 2025", "Mar 2025", "Apr 2025", "May 2025", "Jun 2025", "Jul 2025",
                  "Aug 2025", "Sep 2025", "Oct 2025", "Nov 2025", "Dec 2025", "Jan 2026", "Feb 2026"]


def mom_ttm(path: Path) -> Path:
    """Case 02: 14 months. Sales rise by 100 a month to 2,200, then fall to 1,980; Travel is 0 until Feb 2026."""
    return write_workbook(path, MOM_TTM_MONTHS, [
        (0, "Income", None),                                                                  # 7
        (1, "Product Sales", [1000, 1100, 1200, 1300, 1400, 1500, 1600, 1700, 1800, 1900,
                              2000, 2100, 2200, 1980]),                                       # 8
        (0, "Total Income", "=SUM({c}8:{c}8)"),                                               # 9
        (0, "Cost of Goods Sold", None),                                                      # 10
        (1, "Materials", [400] * 14),                                                         # 11
        (0, "Total Cost of Goods Sold", "=SUM({c}11:{c}11)"),                                 # 12
        (0, "Operating Expenses", None),                                                      # 13
        (1, "Rent", [500] * 14),                                                              # 14
        (1, "Travel", [0] * 13 + [150]),                                                      # 15
        (0, "Total Operating Expenses", "=SUM({c}14:{c}15)"),                                 # 16
    ])


def _simple(sales: list, months: list[str], materials=None, rent=None) -> list[tuple]:
    n = len(months)
    return [
        (0, "Income", None),                                        # 7
        (1, "Product Sales", sales),                                # 8
        (0, "Total Income", "=SUM({c}8:{c}8)"),                     # 9
        (0, "Cost of Goods Sold", None),                            # 10
        (1, "Materials", materials or [300] * n),                   # 11
        (0, "Total Cost of Goods Sold", "=SUM({c}11:{c}11)"),       # 12
        (0, "Operating Expenses", None),                            # 13
        (1, "Rent", rent or [200] * n),                             # 14
        (0, "Total Operating Expenses", "=SUM({c}14:{c}14)"),       # 15
    ]


def missing_month(path: Path) -> Path:
    """Case 03: March 2026 is absent between February and April."""
    months = ["Jan 2026", "Feb 2026", "Apr 2026", "May 2026"]
    return write_workbook(path, months, _simple([1000, 1100, 1300, 1400], months))


def duplicate_period(path: Path) -> Path:
    """Case 04: 'Mar 2026' heads two columns (F and G) with different figures; no month is missing."""
    months = ["Jan 2026", "Feb 2026", "Mar 2026", "Mar 2026"]
    return write_workbook(path, months, _simple([1000, 1100, 1200, 9999], months))


def unmapped_account(path: Path) -> Path:
    """Case 05: 'Crypto Mining Fees' (row 17) matches no category rule in the case's config."""
    return write_workbook(path, JAN_MAR_2026, [
        (0, "Income", None),                                        # 7
        (1, "Product Sales", [5000, 5200, 5400]),                   # 8
        (0, "Total Income", "=SUM({c}8:{c}8)"),                     # 9
        (0, "Cost of Goods Sold", None),                            # 10
        (1, "Materials", [1000, 1040, 1080]),                       # 11
        (0, "Total Cost of Goods Sold", "=SUM({c}11:{c}11)"),       # 12
        (0, "Operating Expenses", None),                            # 13
        (1, "Salaries", [2000, 2000, 2000]),                        # 14
        (1, "Rent", [800, 800, 800]),                               # 15
        (1, "Interest Expense", [50, 50, 50]),                      # 16
        (1, "Crypto Mining Fees", [0, 0, 75]),                      # 17
        (0, "Total Operating Expenses", "=SUM({c}14:{c}17)"),       # 18
    ])


def duplicates_and_subtotals(path: Path) -> Path:
    """Case 06: a nested Payroll group with its own subtotal, 'Software' twice, and a stray
    'Total Marketing' row that closes no group. Only the five accounts may be summed."""
    return write_workbook(path, JAN_MAR_2026, [
        (0, "Income", None),                                        # 7
        (1, "Product Sales", [8000, 8000, 9000]),                   # 8
        (0, "Total Income", "=SUM({c}8:{c}8)"),                     # 9
        (0, "Cost of Goods Sold", None),                            # 10
        (1, "Materials", [2000, 2000, 2500]),                       # 11
        (0, "Total Cost of Goods Sold", "=SUM({c}11:{c}11)"),       # 12
        (0, "Operating Expenses", None),                            # 13
        (1, "Payroll", None),                                       # 14
        (2, "Wages", [3000, 3000, 3000]),                           # 15
        (2, "Payroll Taxes", [300, 300, 300]),                      # 16
        (1, "Total Payroll", "=SUM({c}15:{c}16)"),                  # 17
        (1, "Rent", [1000, 1000, 1000]),                            # 18
        (1, "Software", [100, 100, 100]),                           # 19
        (1, "Software", [50, 50, 60]),                              # 20
        (1, "Total Marketing", [999, 999, 999]),                    # 21
        (0, "Total Operating Expenses", "={c}17+{c}18+{c}19+{c}20"),  # 22
    ])


def invalid_required_value(path: Path) -> Path:
    """Case 07: text 'n/a' in Product Sales for March 2026 (cell F8). January and February are clean."""
    return write_workbook(path, JAN_MAR_2026, [
        (0, "Income", None),                                        # 7
        (1, "Product Sales", [1000, 1100, "n/a"]),                  # 8
        (1, "Service Revenue", [200, 210, 220]),                    # 9
        (0, "Total Income", "=SUM({c}8:{c}9)"),                     # 10
        (0, "Cost of Goods Sold", None),                            # 11
        (1, "Materials", [300, 330, 360]),                          # 12
        (0, "Total Cost of Goods Sold", "=SUM({c}12:{c}12)"),       # 13
        (0, "Operating Expenses", None),                            # 14
        (1, "Rent", [400, 400, 400]),                               # 15
        (1, "Software", [50, 60, 70]),                              # 16
        (0, "Total Operating Expenses", "=SUM({c}15:{c}16)"),       # 17
    ])


def formula_cache(path: Path) -> Path:
    """Case 08: no cached formula values anywhere; 'Insurance' (row 15) is an account whose cells are
    formulas without cell references (=40+5), so its recalculated value feeds the figures."""
    return write_workbook(path, JAN_MAR_2026, [
        (0, "Income", None),                                        # 7
        (1, "Product Sales", [1000, 1100, 1200]),                   # 8
        (0, "Total Income", "=SUM({c}8:{c}8)"),                     # 9
        (0, "Cost of Goods Sold", None),                            # 10
        (1, "Materials", [300, 330, 360]),                          # 11
        (0, "Total Cost of Goods Sold", "=SUM({c}11:{c}11)"),       # 12
        (0, "Operating Expenses", None),                            # 13
        (1, "Rent", [400, 400, 400]),                               # 14
        (1, "Insurance", "=40+5"),                                  # 15
        (0, "Total Operating Expenses", "=SUM({c}14:{c}15)"),       # 16
    ])


def subtotal_mismatch(path: Path) -> Path:
    """Case 09: a hard-coded 'Total Income' that disagrees with its account in March (1,250 vs 1,200)."""
    rows = _simple([1000, 1100, 1200], JAN_MAR_2026)
    rows[2] = (0, "Total Income", [1000, 1100, 1250])               # 9
    return write_workbook(path, JAN_MAR_2026, rows)


FIXTURES = {f.__name__: f for f in (normal_month, mom_ttm, missing_month, duplicate_period, unmapped_account,
                                    duplicates_and_subtotals, invalid_required_value, formula_cache, subtotal_mismatch)}
