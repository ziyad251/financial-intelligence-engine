"""Workbook structure handling and reconciliation checks."""

from decimal import Decimal

import openpyxl
import pytest

from financial_analysis.config import Config
from financial_analysis.formula import FormulaError, FormulaEvaluator
from financial_analysis.issues import AnalysisError
from financial_analysis.loader import DERIVED, LEAF, SUBTOTAL, load_workbook
from financial_analysis.periods import Period, parse_period
from financial_analysis.pipeline import analyze

from helpers import make_workbook, months_from

SIMPLE = [
    (0, "Income", None),
    (1, "Sales", [100, 110, 120]),
    (0, "Total Income", "=SUM({c}8:{c}8)"),
    (0, "Operating Expenses", None),
    (1, "Rent", [50, 50, 50]),
    (0, "Total Operating Expenses", "=SUM({c}11:{c}11)"),
    (1, "Net Income", "={c}9-{c}12"),
]


@pytest.mark.parametrize("text,expected", [
    ("Jun 2026", Period(2026, 6)), ("June 2026", Period(2026, 6)), ("30 Jun 2026", Period(2026, 6)),
    ("2026-06", Period(2026, 6)), ("Jun-26", Period(2026, 6)), ("June 30, 2026", Period(2026, 6)),
    ("Total", None), ("Jan - Jun 2026", None), ("For the month ended June 30, 2026", None), ("Account", None),
])
def test_parse_period(text, expected):
    assert parse_period(text) == expected


def test_row_classification(tmp_path):
    wb = load_workbook(make_workbook(tmp_path / "s.xlsx", ["Jan 2026", "Feb 2026", "Mar 2026"], SIMPLE), Config.load())
    kinds = {r.label: r.kind for r in wb.income_statement.rows}
    assert kinds == {"Sales": LEAF, "Total Income": SUBTOTAL, "Rent": LEAF,
                     "Total Operating Expenses": SUBTOTAL, "Net Income": DERIVED}
    assert [g.label for g in wb.income_statement.groups] == ["Income", "Operating Expenses"]


def test_reverse_chronological_columns_are_sorted(tmp_path):
    rev = [(i, l, list(reversed(v)) if isinstance(v, list) else v) for i, l, v in SIMPLE]
    r = analyze(make_workbook(tmp_path / "r.xlsx", ["Mar 2026", "Feb 2026", "Jan 2026"], rev))
    assert r["periods"]["available"] == ["2026-01", "2026-02", "2026-03"]
    assert r["income_statement"]["monthly"]["revenue"]["values"]["2026-01"] == 100.0


def test_missing_period_is_flagged_and_not_bridged(tmp_path):
    r = analyze(make_workbook(tmp_path / "g.xlsx", ["Jan 2026", "Feb 2026", "Apr 2026"], SIMPLE))
    assert r["periods"]["missing_periods"] == ["2026-03"]
    assert r["status"] == "pass_with_warnings"
    rev = next(m for m in r["analysis_period"]["metrics"] if m["metric"] == "revenue")
    assert rev["mom"]["comparison_period"] == "2026-03" and rev["mom"]["abs_change"] is None
    assert rev["mom"]["note"] == "comparison_period_not_in_data"


def test_duplicate_period_is_an_error(tmp_path):
    r = analyze(make_workbook(tmp_path / "d.xlsx", ["Jan 2026", "Feb 2026", "Feb 2026"], SIMPLE))
    assert r["status"] == "fail"
    assert any(i["code"] == "duplicate_period" for i in r["validation"]["issues"])


def test_total_column_is_ignored_not_summed(tmp_path):
    path = make_workbook(tmp_path / "t.xlsx", ["Jan 2026", "Feb 2026", "Total"],
                         [(i, l, [v[0], v[1], v[0] + v[1]] if isinstance(v, list) else v) for i, l, v in SIMPLE])
    r = analyze(path)
    assert r["periods"]["available"] == ["2026-01", "2026-02"]
    assert any(i["code"] == "ignored_column" for i in r["validation"]["issues"])


def test_missing_revenue_section_raises(tmp_path):
    rows = [(0, "Operating Expenses", None), (1, "Rent", [1, 2]), (0, "Total Operating Expenses", "=SUM({c}8:{c}8)")]
    with pytest.raises(AnalysisError):
        analyze(make_workbook(tmp_path / "n.xlsx", ["Jan 2026", "Feb 2026"], rows))


def test_unmapped_section_is_flagged(tmp_path):
    rows = SIMPLE[:6] + [(0, "Mystery Items", None), (1, "Widget", [5, 5, 5]), (0, "Total Mystery Items", "=SUM({c}14:{c}14)")]
    r = analyze(make_workbook(tmp_path / "u.xlsx", ["Jan 2026", "Feb 2026", "Mar 2026"], rows))
    issue = next(i for i in r["validation"]["issues"] if i["code"] == "unmapped_accounts")
    assert issue["details"]["accounts"] == ["Mystery Items > Widget"]


def test_formula_without_cache_is_evaluated(tmp_path):
    wb = load_workbook(make_workbook(tmp_path / "f.xlsx", ["Jan 2026", "Feb 2026", "Mar 2026"], SIMPLE), Config.load())
    total = next(r for r in wb.income_statement.rows if r.label == "Total Income")
    assert total.values[Period(2026, 2)] == Decimal(110)
    assert any(i.code == "formula_evaluated" for i in wb.issues)


def test_formula_evaluator_rejects_unsupported_functions():
    ev = FormulaEvaluator(lambda c, r: 1)
    assert ev.evaluate("=SUM(A1:A3)+A4*2") == 5
    with pytest.raises(FormulaError):
        ev.evaluate("=VLOOKUP(A1,B1:C3,2)")


def _bs(cash, ap, equity):
    return [
        (0, "Assets", None), (1, "Cash", cash), (0, "Total Assets", "=SUM({c}8:{c}8)"),
        (0, "Liabilities", None), (1, "Accounts Payable", ap), (0, "Total Liabilities", "=SUM({c}11:{c}11)"),
        (0, "Equity", None), (1, "Retained Earnings", [0, 0, 0]), (1, "Current Year Earnings", equity),
        (0, "Total Equity", "=SUM({c}14:{c}15)"),
    ]


def test_balance_sheet_out_of_balance_is_an_error(tmp_path):
    # Jan: 90 = 40 + 50 and Mar: 220 = 40 + 180 balance; Feb: 200 != 40 + 110
    r = analyze(make_workbook(tmp_path / "b.xlsx", ["Jan 2026", "Feb 2026", "Mar 2026"], SIMPLE,
                              _bs([90, 200, 220], [40, 40, 40], [50, 110, 180])))
    assert r["status"] == "fail"
    issue = next(i for i in r["validation"]["issues"] if i["code"] == "balance_sheet_out_of_balance")
    assert [m["period"] for m in issue["details"]["mismatches"]] == ["2026-02"]


def test_earnings_rollforward(tmp_path):
    # Net income: 100-50=50, 110-50=60, 120-50=70. Current-year earnings: 50, 110, 180 (Jan = NI, then cumulative).
    good = analyze(make_workbook(tmp_path / "g.xlsx", ["Jan 2026", "Feb 2026", "Mar 2026"], SIMPLE,
                                 _bs([90, 150, 220], [40, 40, 40], [50, 110, 180])))
    check = next(c for c in good["validation"]["checks"] if c["check"] == "earnings_rollforward")
    assert check["status"] == "pass" and check["periods_compared"] == 3
    bad = analyze(make_workbook(tmp_path / "x.xlsx", ["Jan 2026", "Feb 2026", "Mar 2026"], SIMPLE,
                                _bs([90, 150, 230], [40, 40, 40], [50, 110, 190])))
    issue = next(i for i in bad["validation"]["issues"] if i["code"] == "earnings_rollforward_mismatch")
    assert [m["period"] for m in issue["details"]["mismatches"]] == ["2026-03"]


def test_highlight_rules(tmp_path):
    # Materiality = max(5000, 2% of revenue). Revenue 100,000 -> 5,000.
    months = months_from(2026, 1, 2)
    rows = [
        (0, "Income", None), (1, "Sales", [100000, 100000]), (0, "Total Income", "=SUM({c}8:{c}8)"),
        (0, "Operating Expenses", None),
        (1, "Travel", [1000, 7000]),     # +6,000 and +600%: flagged
        (1, "Payroll", [60000, 64000]),  # +4,000: below materiality
        (1, "Hosting", [40000, 46000]),  # +6,000 but +15% < 20%: not flagged
        (0, "Total Operating Expenses", "=SUM({c}11:{c}13)"),
    ]
    r = analyze(make_workbook(tmp_path / "h.xlsx", months, rows))
    flagged = {(h["key"], h["type"]) for h in r["highlights"]}
    assert ("category:Travel", "mom_change") in flagged
    assert not any("Payroll" in k or "Hosting" in k for k, _ in flagged)
    # Opex total +16,000 >= 5,000: headline flagged; operating income -16,000 flagged
    assert ("opex_total", "mom_change") in flagged and ("operating_income", "mom_change") in flagged
    assert r["periods"]["materiality_threshold"] == 5000.0
