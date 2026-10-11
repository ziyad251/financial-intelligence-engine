"""Build small financial-statement workbooks for tests.

Rows are (indent, label, values):
  values None             -> section/group header (no values)
  values list             -> one entry per month: number, None (blank) or text
  values "=SUM({c}8:{c}9)" -> formula template; {c} is replaced by each month's column letter
  values "@...", e.g. "@40" -> hard-coded (non-formula) number repeated in every month

openpyxl does not compute formulas, so saved files have no cached values. The loader
evaluates them with its own formula evaluator, which the tests exercise as a result.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HEADER_ROW = 5
FIRST_DATA_ROW = 7


def write_sheet(ws, months: list[str], rows: list[tuple], label_cols: int = 2) -> dict[str, int]:
    """Write rows starting at FIRST_DATA_ROW; return {label: sheet row}."""
    ws.cell(1, 1, "Test statement")
    ws.cell(HEADER_ROW, label_cols, "Account")
    first = label_cols + 1
    for i, m in enumerate(months):
        ws.cell(HEADER_ROW, first + i, m)
    where: dict[str, int] = {}
    r = FIRST_DATA_ROW
    for indent, label, values in rows:
        ws.cell(r, 1 + indent, label)
        where[label] = r
        for i in range(len(months)):
            col = first + i
            if values is None:
                continue
            if isinstance(values, str) and values.startswith("="):
                ws.cell(r, col, values.format(c=get_column_letter(col)))
            elif isinstance(values, str) and values.startswith("@"):
                ws.cell(r, col, float(values[1:]))
            else:
                v = values[i]
                if v is not None:
                    ws.cell(r, col, v)
        r += 1
    return where


def make_workbook(path: Path, months: list[str], is_rows: list[tuple], bs_rows: list[tuple] | None = None,
                  bs_months: list[str] | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    write_sheet(ws, months, is_rows)
    if bs_rows is not None:
        write_sheet(wb.create_sheet("Balance Sheet"), bs_months or months, bs_rows)
    wb.save(path)
    return path


def months_from(year: int, month: int, n: int) -> list[str]:
    import calendar
    out = []
    for i in range(n):
        y, m = divmod(month - 1 + i, 12)
        out.append(f"{calendar.month_abbr[m + 1]} {year + y}")
    return out
