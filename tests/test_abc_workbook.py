"""Integration test on the supplied ABC Inc workbook.

The expected values are not taken from this pipeline. They come from the totals Excel
cached in the workbook itself (Total Income, Net Income and so on, read here
directly with openpyxl) or from sums of those cached totals. The test is skipped if
the workbook is not present.
"""

from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from financial_analysis.config import Config
from financial_analysis.pipeline import analyze

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "ABC Inc - Financial Statements (1) (1).xlsx"
CONFIG = ROOT / "config" / "companies" / "abc-inc.json"

pytestmark = pytest.mark.skipif(not WORKBOOK.exists(), reason="ABC workbook not available")

MONTHS = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
          "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}


@pytest.fixture(scope="module")
def excel_totals():
    """{row label: {YYYY-MM: Decimal}} for total rows, from Excel's cached values."""
    ws = openpyxl.load_workbook(WORKBOOK, data_only=True)["Income Statement"]
    header = {c: ws.cell(5, c).value for c in range(3, ws.max_column + 1) if ws.cell(5, c).value}
    keys = {c: f"{h.split()[1]}-{MONTHS[h.split()[0]]:02d}" for c, h in header.items()}
    out = {}
    for r in range(6, ws.max_row + 1):
        label = ws.cell(r, 1).value or ws.cell(r, 2).value
        if label in ("Total Income", "Total Cost of Goods Sold", "Gross Profit", "Total Operating Expenses", "Net Income"):
            out[label] = {keys[c]: Decimal(str(ws.cell(r, c).value)) for c in keys}
    return out


@pytest.fixture(scope="module")
def result():
    return analyze(WORKBOOK, Config.load(CONFIG), "2026-06")


def _q(d: Decimal) -> float:
    return float(d.quantize(Decimal("0.01")))


def test_status_and_periods(result):
    assert result["status"] == "pass"
    assert result["periods"]["available"][0] == "2025-01" and result["periods"]["available"][-1] == "2026-06"
    assert len(result["periods"]["available"]) == 18
    assert result["periods"]["missing_periods"] == []


@pytest.mark.parametrize("metric,label", [("revenue", "Total Income"), ("cogs", "Total Cost of Goods Sold"),
                                          ("gross_profit", "Gross Profit"), ("opex_total", "Total Operating Expenses"),
                                          ("net_income", "Net Income")])
def test_every_month_matches_excel(result, excel_totals, metric, label):
    calculated = result["income_statement"]["monthly"][metric]["values"]
    assert calculated == {k: _q(v) for k, v in excel_totals[label].items()}


def test_ttm_revenue_matches_sum_of_excel_months(result, excel_totals):
    months = [f"2025-{m:02d}" for m in range(7, 13)] + [f"2026-{m:02d}" for m in range(1, 7)]
    expected = sum(excel_totals["Total Income"][k] for k in months)
    assert result["income_statement"]["ttm"]["revenue"]["values"]["2026-06"] == _q(expected)
    assert result["income_statement"]["ttm"]["revenue"]["values"]["2025-11"] is None


def test_known_events_are_surfaced(result):
    outliers = {(h["label"], h["period"]) for h in result["history_outliers"]}
    assert ("Income Tax Benefit", "2025-12") in outliers
    assert ("Stock-Based Compensation Expense", "2025-12") in outliers
    rev_mom = [h for h in result["highlights"] if h["key"] == "revenue" and h["type"] == "mom_change"]
    assert rev_mom and rev_mom[0]["direction"] == "decrease"


def test_every_opex_account_is_categorised(result):
    names = [c["name"] for c in result["income_statement"]["opex_categories"]]
    assert "Other operating expenses" not in names


def test_reconciliations_pass(result):
    statuses = {c["check"]: c["status"] for c in result["validation"]["checks"]}
    assert statuses["balance_sheet_balances"] == "pass"
    assert statuses["earnings_rollforward"] == "pass"
    assert statuses["subtotal:Income Statement:Operating Expenses"] == "pass"
    assert statuses["subtotal:Income Statement:S & M Payroll"] == "pass"
